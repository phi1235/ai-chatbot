"""Parse uploaded files (PDF, DOCX, MD, TXT) thành plain text để chunk + embed.

Hỗ trợ:
    .pdf  - pypdf
    .docx - python-docx
    .md / .markdown / .txt / không extension - đọc utf-8 trực tiếp

Trả về string text đã clean. Caller chịu trách nhiệm hash → doc_id, chunk, embed.
"""
from __future__ import annotations

import io
import re
from pathlib import Path

# Giới hạn để tránh OOM với file siêu lớn
MAX_FILE_SIZE_BYTES = 20 * 1024 * 1024  # 20MB


class FileParseError(Exception):
    pass


def _clean_text(text: str) -> str:
    """Bỏ khoảng trắng thừa, normalize line breaks."""
    # Trim trailing spaces mỗi dòng
    text = re.sub(r"[ \t]+\n", "\n", text)
    # Gộp 3+ newlines thành 2
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def parse_pdf(data: bytes) -> str:
    from pypdf import PdfReader

    try:
        reader = PdfReader(io.BytesIO(data))
    except Exception as exc:
        raise FileParseError(f"Không đọc được PDF: {exc}") from exc

    if reader.is_encrypted:
        try:
            # Thử password rỗng (PDF protect read-only)
            reader.decrypt("")
        except Exception as exc:
            raise FileParseError("PDF được mã hoá, cần password.") from exc

    pages: list[str] = []
    for i, page in enumerate(reader.pages, start=1):
        try:
            text = page.extract_text() or ""
        except Exception:
            text = ""
        text = text.strip()
        if text:
            pages.append(f"# Page {i}\n\n{text}")

    if not pages:
        raise FileParseError("PDF không có text trích xuất được (có thể là scan ảnh).")
    return _clean_text("\n\n".join(pages))


def parse_docx(data: bytes) -> str:
    from docx import Document

    try:
        doc = Document(io.BytesIO(data))
    except Exception as exc:
        raise FileParseError(f"Không đọc được DOCX: {exc}") from exc

    parts: list[str] = []
    for para in doc.paragraphs:
        text = (para.text or "").strip()
        if not text:
            continue
        # Heading style → giữ heading marker để chunker tách section
        style_name = (para.style.name or "").lower() if para.style else ""
        if style_name.startswith("heading"):
            parts.append(f"## {text}")
        else:
            parts.append(text)

    # Bảng (table) - extract text từng cell, ngăn cách bằng |
    for table in doc.tables:
        for row in table.rows:
            cells = [(cell.text or "").strip() for cell in row.cells]
            cells = [c for c in cells if c]
            if cells:
                parts.append(" | ".join(cells))

    if not parts:
        raise FileParseError("DOCX không có nội dung text.")
    return _clean_text("\n\n".join(parts))


def parse_text(data: bytes) -> str:
    """Đọc text/markdown/plain - decode utf-8 (fallback latin-1)."""
    try:
        text = data.decode("utf-8")
    except UnicodeDecodeError:
        text = data.decode("latin-1", errors="replace")
    return _clean_text(text)


def parse_file(filename: str, data: bytes) -> str:
    """Dispatch parser theo extension."""
    if len(data) > MAX_FILE_SIZE_BYTES:
        raise FileParseError(f"File quá lớn (>{MAX_FILE_SIZE_BYTES // 1024 // 1024}MB).")
    if not data:
        raise FileParseError("File rỗng.")

    ext = Path(filename).suffix.lower()
    if ext == ".pdf":
        return parse_pdf(data)
    if ext == ".docx":
        return parse_docx(data)
    if ext in (".md", ".markdown", ".txt", ".rst", ""):
        return parse_text(data)
    raise FileParseError(f"Định dạng không hỗ trợ: {ext or '(không có ext)'}")

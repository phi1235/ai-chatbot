"""Tests cho processor.file_parser."""
from __future__ import annotations

import io

import pytest

from processor.file_parser import (
    FileParseError,
    parse_docx,
    parse_file,
    parse_pdf,
    parse_text,
)


# ─── parse_text ──────────────────────────────────────────────────────────────
def test_parse_text_utf8():
    data = "Đây là tiếng Việt có dấu.\n\nDòng 2.".encode()
    assert parse_text(data) == "Đây là tiếng Việt có dấu.\n\nDòng 2."


def test_parse_text_latin1_fallback():
    """File encoding lạ vẫn parse được bằng latin-1 fallback."""
    # Bytes không phải utf-8 hợp lệ
    data = b"\xc3\x28 broken utf8"  # invalid utf-8 sequence
    out = parse_text(data)
    assert out  # không raise


def test_parse_text_normalizes_whitespace():
    data = b"line1   \n\n\n\nline2  \n"
    out = parse_text(data)
    assert out == "line1\n\nline2"


# ─── parse_pdf ───────────────────────────────────────────────────────────────
def test_parse_pdf_extracts_text():
    """Tạo PDF tối giản với pypdf rồi đọc lại."""
    from pypdf import PdfWriter

    writer = PdfWriter()
    writer.add_blank_page(width=612, height=792)
    buf = io.BytesIO()
    writer.write(buf)
    pdf_bytes = buf.getvalue()

    # PDF blank → no text extracted → raise
    with pytest.raises(FileParseError, match="không có text"):
        parse_pdf(pdf_bytes)


def test_parse_pdf_invalid_data():
    with pytest.raises(FileParseError, match="Không đọc được"):
        parse_pdf(b"not a pdf at all")


# ─── parse_docx ──────────────────────────────────────────────────────────────
def test_parse_docx_basic():
    from docx import Document

    doc = Document()
    doc.add_heading("Tiêu đề chính", level=1)
    doc.add_paragraph("Đoạn văn thứ nhất.")
    doc.add_paragraph("Đoạn văn thứ hai.")
    buf = io.BytesIO()
    doc.save(buf)

    out = parse_docx(buf.getvalue())
    assert "Tiêu đề chính" in out
    assert "## " in out  # heading được mark
    assert "Đoạn văn thứ nhất." in out
    assert "Đoạn văn thứ hai." in out


def test_parse_docx_with_table():
    from docx import Document

    doc = Document()
    doc.add_paragraph("Trước bảng.")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "A"
    table.cell(0, 1).text = "B"
    table.cell(1, 0).text = "C"
    table.cell(1, 1).text = "D"
    buf = io.BytesIO()
    doc.save(buf)

    out = parse_docx(buf.getvalue())
    assert "A | B" in out
    assert "C | D" in out


def test_parse_docx_invalid():
    with pytest.raises(FileParseError):
        parse_docx(b"not docx")


# ─── parse_file dispatcher ───────────────────────────────────────────────────
def test_parse_file_dispatch_txt():
    out = parse_file("notes.txt", b"hello world")
    assert out == "hello world"


def test_parse_file_dispatch_md():
    out = parse_file("readme.md", b"# Title\n\ncontent")
    assert "# Title" in out
    assert "content" in out


def test_parse_file_unsupported_extension():
    with pytest.raises(FileParseError, match="không hỗ trợ"):
        parse_file("image.png", b"binary png data")


def test_parse_file_too_large():
    big = b"x" * (21 * 1024 * 1024)  # 21MB > 20MB limit
    with pytest.raises(FileParseError, match="quá lớn"):
        parse_file("big.txt", big)


def test_parse_file_empty():
    with pytest.raises(FileParseError, match="rỗng"):
        parse_file("empty.txt", b"")
